"""Workbook parser and OHLC validity. Adapted from energy-vol-targeting/tests/test_load.py and test_quality.py."""
import datetime as dt

import numpy as np
import openpyxl
import pandas as pd

from etr.returns import ohlc_valid, read_workbook, wide


def _make_workbook(path):
    """Two sheets mimicking the Bloomberg layout, plus a README sheet that must be skipped."""
    wb = openpyxl.Workbook()
    wb.active.title = "README"
    wb.active.append(["not data"])
    ws = wb.create_sheet("A_GasCurves")
    ws.append(["[Tier B] title"])
    ws.append(["description"])
    ws.append(["Date", "PX_LAST", None, "Date", "PX_LAST"])
    ws.append(["TZT1 Comdty", None, None, "TZT2 Comdty", None])
    ws.append(["TTF M1", None, None, "TTF M2", None])
    ws.append(["EUR/MWh | conf H", None, None, "EUR/MWh | conf H", None])
    d = [dt.datetime(2020, 1, i) for i in (1, 2, 3)]
    ws.append([d[0], "", None, d[0], ""])           # holiday: blank, must stay NaN
    ws.append([d[1], 10.5, None, d[1], 11.0])
    ws.append([d[2], 10.0, None, d[2], "#N/A N/A"])  # Bloomberg error string -> NaN
    ws2 = wb.create_sheet("E_OHLC_Vol")
    ws2.append(["t"]); ws2.append(["d"])
    ws2.append(["Date", "PX_HIGH", "PX_LAST"])
    ws2.append(["TZT1 Comdty", None, None])
    ws2.append(["TTF M1 OHLC", None, None])
    ws2.append(["EUR/MWh", None, None])
    ws2.append([d[1], 10.9, 99.0])  # conflicting PX_LAST: curve sheet should win
    wb.save(path)


def test_read_workbook_layout(tmp_path):
    p = tmp_path / "wb.xlsx"
    _make_workbook(p)
    long = read_workbook(p)
    assert set(long["sheet"]) == {"A_GasCurves", "E_OHLC_Vol"}
    m1 = long[(long.sheet == "A_GasCurves") & (long.ticker == "TZT1 Comdty")].set_index("date")["value"]
    assert np.isnan(m1[pd.Timestamp("2020-01-01")])   # blank stays blank, not 0, not filled
    assert m1[pd.Timestamp("2020-01-02")] == 10.5
    m2 = long[long.ticker == "TZT2 Comdty"].set_index("date")["value"]
    assert np.isnan(m2[pd.Timestamp("2020-01-03")])
    assert set(long[long.sheet == "E_OHLC_Vol"]["field"]) == {"PX_HIGH", "PX_LAST"}


def test_wide_prefers_curve_sheet_and_drops_empty_rows(tmp_path):
    p = tmp_path / "wb.xlsx"
    _make_workbook(p)
    w = wide(read_workbook(p), "PX_LAST", ["TZT1 Comdty", "TZT2 Comdty"])
    assert w.loc["2020-01-02", "TZT1 Comdty"] == 10.5     # not the 99.0 from E_OHLC_Vol
    assert pd.Timestamp("2020-01-01") not in w.index       # all-blank row dropped
    assert list(w.columns) == ["TZT1 Comdty", "TZT2 Comdty"]


def test_ohlc_valid():
    s = lambda *v: pd.Series(v, dtype=float)  # noqa: E731
    o = s(10, 10, 10, 10, np.nan, 10)
    h = s(11, 11, 11, 9, 11, 10)
    l = s(9, 9, 9, 8, 9, 10)
    c = s(10.5, 11.5, 8.5, 8.5, 10, 10)
    #       ok   close>high close<low high<open missing  flat bar (ok)
    assert ohlc_valid(o, h, l, c).tolist() == [True, False, False, False, False, True]
